#!/usr/bin/env python3
"""learning_loop.py -- close the loop between backtests, the journal,
and the next round of simulated experiments.

Usage:
    python3 analysis/learning_loop.py \
        [--results backtests/results.jsonl] \
        [--journal Trading_Journal.xlsx] \
        [--out reports/learning_report.md]

Defaults are relative to the DayTrading root (resolved from this
script's location). Reads the journal, never writes it.

OUTPUT IS ADVISORY ONLY. No brokerage access; nothing here places,
stages, or recommends orders. Every report ends with the exact line:
"Recommendations are hypotheses to test in simulation - not
live-trading instructions."
"""

import argparse
import json
import os
import re
import sys
from datetime import datetime

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_RESULTS = os.path.join(ROOT, "backtests", "results.jsonl")
DEFAULT_SESSIONS = os.path.join(ROOT, "backtests", "sessions.jsonl")
DEFAULT_JOURNAL = os.path.join(ROOT, "Trading_Journal.xlsx")
DEFAULT_OUT = os.path.join(ROOT, "reports", "learning_report.md")

SESSION_VERDICTS = ("CORRECT_TRADE", "CORRECT_NO_TRADE", "MISSED_SIGNAL",
                    "OFF_PLAN_TRADE")
SESSION_ADHERENCE_GATE_PCT = 80.0  # mirrors strategy.md section 13
TRADING_DAYS_PER_WEEK = 5
ORB_TIEBREAK_WARNING = ("ORB same-bar priority remains UNVALIDATED "
                        "(6-trade sample).")

WATCHLIST = ("SPY", "QQQ", "XLF", "XLE", "IWM")
EXAMPLE_SCREENSHOT = "screenshots/2026-07-13-xlf.png"
FOOTER = ("Recommendations are hypotheses to test in simulation - "
          "not live-trading instructions.")
DAYS_PER_MONTH = 30.44

# Journal Trade Log columns (1-indexed fallbacks if headers are missing).
JOURNAL_COLS = {
    "Date": 1, "Time (CT)": 2, "Ticker": 3, "Direction": 4, "Setup": 5,
    "Screenshot": 30, "Realized R": 33, "Highest chart R": 34,
    "MFE ($)": 35, "MAE ($)": 36, "Exit reason": 37, "Followed plan": 38,
    "Chased": 39, "Widened stop": 40, "Lesson": 41,
}


# ---------------------------------------------------------------- helpers

def _num(v):
    """Coerce a cell/metric to float, else None."""
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        try:
            return float(v.strip())
        except ValueError:
            return None
    return None


def _fmt(v, digits=2):
    return "n/a" if v is None else ("%.*f" % (digits, v))


def _yn(v):
    """Normalize a Y/N-ish cell to 'Y', 'N', or None."""
    if v is None:
        return None
    s = str(v).strip().upper()
    if s in ("Y", "YES", "TRUE", "1"):
        return "Y"
    if s in ("N", "NO", "FALSE", "0"):
        return "N"
    return None


# ---------------------------------------------------------- backtest side

def load_results(path):
    """Return (records, warnings). Malformed lines are skipped and noted."""
    records, warnings = [], []
    if not os.path.exists(path):
        warnings.append("results file not found: %s (treated as empty)" % path)
        return records, warnings
    with open(path, "r") as fh:
        for lineno, line in enumerate(fh, 1):
            stripped = line.strip()
            if not stripped:
                continue
            try:
                obj = json.loads(stripped)
            except ValueError as exc:
                warnings.append("line %d: not valid JSON (%s): %s"
                                % (lineno, exc, stripped[:80]))
                continue
            if not isinstance(obj, dict):
                warnings.append("line %d: not a JSON object, skipped" % lineno)
                continue
            for key in ("run_id", "symbol", "level_source", "config",
                        "metrics"):
                if key not in obj:
                    warnings.append("line %d: missing '%s', skipped"
                                    % (lineno, key))
                    break
            else:
                records.append(obj)
    return records, warnings


def _cfg(rec, key):
    cfg = rec.get("config") or {}
    return cfg.get(key)


def _metric(rec, key):
    metrics = rec.get("metrics") or {}
    return metrics.get(key)


def _trades(rec):
    v = _metric(rec, "trades")
    return int(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def config_signature(rec):
    return (rec.get("symbol"), rec.get("level_source"),
            _cfg(rec, "entry_mode"), _cfg(rec, "exit_mode"))


def period_months(rec):
    period = rec.get("period") or {}
    try:
        d0 = datetime.strptime(str(period.get("from")), "%Y-%m-%d")
        d1 = datetime.strptime(str(period.get("to")), "%Y-%m-%d")
    except (ValueError, TypeError):
        return None
    days = (d1 - d0).days
    return max(days, 1) / DAYS_PER_MONTH


def aggregate(records, keyfunc):
    """Group runs; sum trades/wins/net_r per group. Returns
    {key: {"runs": n, "trades": t|None, "wins": w|None, "net_r": r|None,
           "wr_trades": trades counted toward weighted win rate}}."""
    groups = {}
    for rec in records:
        key = keyfunc(rec)
        g = groups.setdefault(key, {"runs": 0, "trades": 0, "wins": 0,
                                    "net_r": 0.0, "wr_trades": 0,
                                    "has_trades": False, "has_net": False})
        g["runs"] += 1
        t = _trades(rec)
        if t is not None:
            g["trades"] += t
            g["has_trades"] = True
        w = _num(_metric(rec, "wins"))
        if t is not None and w is not None:
            g["wins"] += w
            g["wr_trades"] += t
        nr = _num(_metric(rec, "net_r"))
        if nr is not None:
            g["net_r"] += nr
            g["has_net"] = True
    for g in groups.values():
        if not g["has_trades"]:
            g["trades"] = None
        if not g["has_net"]:
            g["net_r"] = None
    return groups


def weighted_win_rate(group):
    if group["wr_trades"] > 0:
        return 100.0 * group["wins"] / group["wr_trades"]
    return None


# ---------------------------------------------------------- sessions side
# (backtests/sessions.jsonl -- one row per trading DAY, recorded by
# analysis/record_session.py. Closes the blind spot where a legitimate
# no-trade day, with the engine correctly producing zero QUALIFIED events,
# used to leave no row anywhere in the learning loop.)

def load_sessions(path):
    """Return (records, warnings). Malformed lines are skipped and noted."""
    records, warnings = [], []
    if not os.path.exists(path):
        warnings.append("sessions file not found: %s (treated as empty)" % path)
        return records, warnings
    with open(path, "r") as fh:
        for lineno, line in enumerate(fh, 1):
            stripped = line.strip()
            if not stripped:
                continue
            try:
                obj = json.loads(stripped)
            except ValueError as exc:
                warnings.append("sessions line %d: not valid JSON (%s): %s"
                                % (lineno, exc, stripped[:80]))
                continue
            if not isinstance(obj, dict):
                warnings.append("sessions line %d: not a JSON object, skipped"
                                % lineno)
                continue
            for key in ("session_date", "engine", "trader", "verdict"):
                if key not in obj:
                    warnings.append("sessions line %d: missing '%s', skipped"
                                    % (lineno, key))
                    break
            else:
                records.append(obj)
    return records, warnings


def session_confusion_matrix(records):
    """{'CORRECT_TRADE': n, 'CORRECT_NO_TRADE': n, 'MISSED_SIGNAL': n,
    'OFF_PLAN_TRADE': n}."""
    matrix = {v: 0 for v in SESSION_VERDICTS}
    for r in records:
        v = r.get("verdict")
        if v in matrix:
            matrix[v] += 1
    return matrix


def session_decision_accuracy_pct(matrix, total):
    """(CORRECT_TRADE + CORRECT_NO_TRADE) / total sessions, as a percentage.
    None if there are no sessions yet."""
    if not total:
        return None
    correct = matrix["CORRECT_TRADE"] + matrix["CORRECT_NO_TRADE"]
    return 100.0 * correct / total


def session_qualified_frequency(records):
    """Returns dict with total_qualified, sessions_with_known_engine,
    per_session, per_week (assuming TRADING_DAYS_PER_WEEK session/week),
    and excluded (sessions recorded via --no-events, engine.qualified
    unknown -- never averaged in, since an absent replay can't attest to
    zero OR nonzero QUALIFIED events)."""
    known = [r for r in records
             if isinstance(r.get("engine"), dict)
             and r["engine"].get("qualified") is not None]
    total_q = sum(r["engine"]["qualified"] for r in known)
    n = len(known)
    per_session = (total_q / float(n)) if n else None
    per_week = (per_session * TRADING_DAYS_PER_WEEK) if per_session is not None else None
    return {
        "total_qualified": total_q,
        "sessions_with_known_engine": n,
        "per_session": per_session,
        "per_week": per_week,
        "excluded": len(records) - n,
    }


def session_total_pre_window_rejects(records):
    """(total, sessions_with_known_count, sessions_excluded_unknown)."""
    vals = [r["engine"].get("pre_window_rejects") for r in records
            if isinstance(r.get("engine"), dict)]
    known = [v for v in vals if v is not None]
    return sum(known), len(known), len(vals) - len(known)


# ----------------------------------------------------------- journal side

def load_journal(path):
    """Read the Trade Log sheet (read-only). Returns a dict:
    {"found": bool, "note": str, "rows": [row dicts], "examples_excluded": n}
    Never writes the workbook."""
    out = {"found": False, "note": "", "rows": [], "examples_excluded": 0}
    if not os.path.exists(path):
        out["note"] = "journal file not found at %s" % path
        return out
    try:
        import openpyxl
    except ImportError:
        out["note"] = "openpyxl not available; journal not read"
        return out
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # corrupt / locked file must not kill the report
        out["note"] = "journal could not be opened (%s)" % exc
        return out
    try:
        if "Trade Log" not in wb.sheetnames:
            out["note"] = "no 'Trade Log' sheet in %s" % path
            return out
        ws = wb["Trade Log"]
        out["found"] = True

        header_idx = dict(JOURNAL_COLS)  # fallbacks
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            out["note"] = "Trade Log sheet is empty"
            return out
        for idx, name in enumerate(header_row, 1):
            if isinstance(name, str) and name.strip() in JOURNAL_COLS:
                header_idx[name.strip()] = idx

        def cell(row, name):
            i = header_idx[name] - 1
            return row[i] if i < len(row) else None

        for row in rows_iter:
            if row is None:
                continue
            ticker = cell(row, "Ticker")
            if ticker is None or not str(ticker).strip():
                continue
            screenshot = cell(row, "Screenshot")
            if (screenshot is not None
                    and str(screenshot).strip() == EXAMPLE_SCREENSHOT):
                out["examples_excluded"] += 1
                continue
            text_blob = " ".join(
                str(cell(row, n) or "") for n in ("Exit reason", "Lesson"))
            out["rows"].append({
                "ticker": str(ticker).strip().upper(),
                "setup": (str(cell(row, "Setup")).strip()
                          if cell(row, "Setup") is not None else "unspecified"),
                "followed": _yn(cell(row, "Followed plan")),
                "chased": _yn(cell(row, "Chased")),
                "widened": _yn(cell(row, "Widened stop")),
                "realized_r": _num(cell(row, "Realized R")),
                "highest_r": _num(cell(row, "Highest chart R")),
                "mfe": _num(cell(row, "MFE ($)")),
                "mae": _num(cell(row, "MAE ($)")),
                "is_sim": bool(re.search(r"\b(sim|replay|drill)\b",
                                         text_blob, re.IGNORECASE)),
            })
    finally:
        wb.close()
    return out


def journal_stats(rows):
    """Execution metrics from real journal rows."""
    stats = {"n": len(rows)}
    scored = [r for r in rows if r["followed"] is not None]
    stats["adherence_n"] = len(scored)
    stats["adherence_pct"] = (
        100.0 * sum(1 for r in scored if r["followed"] == "Y") / len(scored)
        if scored else None)
    stats["chases"] = sum(1 for r in rows if r["chased"] == "Y")
    stats["widened"] = sum(1 for r in rows if r["widened"] == "Y")

    by_setup = {}
    for r in rows:
        if r["realized_r"] is None:
            continue
        by_setup.setdefault(r["setup"], []).append(r["realized_r"])
    stats["r_by_setup"] = {
        k: (sum(v) / len(v), sum(v), len(v)) for k, v in sorted(by_setup.items())}

    utils = [r["realized_r"] / r["highest_r"] for r in rows
             if r["realized_r"] is not None and r["highest_r"] not in (None, 0.0)
             and r["highest_r"] > 0]
    stats["mfe_util_pct"] = 100.0 * sum(utils) / len(utils) if utils else None
    mfes = [r["mfe"] for r in rows if r["mfe"] is not None]
    maes = [r["mae"] for r in rows if r["mae"] is not None]
    stats["avg_mfe"] = sum(mfes) / len(mfes) if mfes else None
    stats["avg_mae"] = sum(maes) / len(maes) if maes else None

    followed_r = [r["realized_r"] for r in rows
                  if r["followed"] == "Y" and r["realized_r"] is not None]
    broken_r = [r["realized_r"] for r in rows
                if r["followed"] == "N" and r["realized_r"] is not None]
    stats["followed_r"] = followed_r
    stats["broken_r"] = broken_r
    return stats


# -------------------------------------------------------- recommendations

def confidence(sample):
    """LOW unless supporting sample >= 20 trades; never above MEDIUM
    from backtest evidence alone."""
    return "MEDIUM" if (sample is not None and sample >= 20) else "LOW"


def build_recommendations(records, journal, jstats):
    """Rule-based recommendations. Returns list of dicts:
    {"text", "evidence", "confidence"}."""
    recs = []

    # (a) config with net_r <= 0 across >= 10 trades -> recommend against.
    by_cfg = aggregate(records, config_signature)
    for key in sorted(by_cfg, key=lambda k: tuple(str(x) for x in k)):
        g = by_cfg[key]
        if g["trades"] is not None and g["trades"] >= 10 \
                and g["net_r"] is not None and g["net_r"] <= 0:
            sym, lvl, entry, exit_mode = key
            recs.append({
                "text": ("Recommend AGAINST the config %s / %s / %s / %s in "
                         "further sims; it is net-negative at sample size."
                         % (sym, lvl, entry, exit_mode)),
                "evidence": ("%d trades across %d run(s), combined net R "
                             "%.2f (<= 0)." % (g["trades"], g["runs"],
                                               g["net_r"])),
                "confidence": confidence(g["trades"]),
            })

    # (b) ORB beats PDH/PDL on the same symbol by > 1R, both >= 10 trades.
    by_sym_lvl = aggregate(
        records, lambda r: (r.get("symbol"), r.get("level_source")))
    symbols = sorted({r.get("symbol") for r in records if r.get("symbol")})
    for sym in symbols:
        orb = by_sym_lvl.get((sym, "ORH/ORL"))
        pd_ = by_sym_lvl.get((sym, "PDH/PDL"))
        if not orb or not pd_:
            continue
        if (orb["trades"] is not None and orb["trades"] >= 10
                and pd_["trades"] is not None and pd_["trades"] >= 10
                and orb["net_r"] is not None and pd_["net_r"] is not None
                and orb["net_r"] - pd_["net_r"] > 1.0):
            recs.append({
                "text": ("Prioritize ORB (ORH/ORL) setups over PDH/PDL on "
                         "%s in the next sim block." % sym),
                "evidence": ("ORH/ORL net R %.2f on %d trades vs PDH/PDL "
                             "net R %.2f on %d trades (edge %.2fR > 1R)."
                             % (orb["net_r"], orb["trades"], pd_["net_r"],
                                pd_["trades"],
                                orb["net_r"] - pd_["net_r"])),
                "confidence": confidence(min(orb["trades"], pd_["trades"])),
            })

    # (c) symbol whose every run fired zero trades.
    for sym in symbols:
        sym_runs = [r for r in records if r.get("symbol") == sym]
        trade_counts = [_trades(r) for r in sym_runs]
        if trade_counts and all(t == 0 for t in trade_counts):
            rvols = sorted({_cfg(r, "rvol_min") for r in sym_runs
                            if _num(_cfg(r, "rvol_min")) is not None})
            rvol_note = (" (e.g. rvol_min %s -> %.1f)"
                         % (rvols[0], float(rvols[0]) - 0.1)) if rvols else ""
            recs.append({
                "text": ("%s produced zero trades in every tested config: "
                         "either drop %s from the active scan or loosen ONE "
                         "named filter as a controlled experiment -- "
                         "suggested single variable: rvol_min%s."
                         % (sym, sym, rvol_note)),
                "evidence": ("0 trades across %d run(s) on %s."
                             % (len(sym_runs), sym)),
                "confidence": "LOW",
            })

    # (e) every config fires < 1 trade/month -> frequency too low to
    # validate the edge; expand the sim regime, do not chase signals.
    rate_runs = []
    for r in records:
        t = _trades(r)
        months = period_months(r)
        if t is not None and months:
            rate_runs.append(t / months)
    if rate_runs and all(rate < 1.0 for rate in rate_runs):
        recs.append({
            "text": ("Every tested config fires < 1 trade/month; the edge "
                     "cannot be validated at this frequency. Expand the sim "
                     "regime (scheduled Bar Replay drills across more "
                     "sessions) rather than loosening filters to chase "
                     "signals."),
            "evidence": ("Max observed rate %.2f trades/month across %d "
                         "run(s) with usable periods."
                         % (max(rate_runs), len(rate_runs))),
            "confidence": "LOW",
        })

    # (d) adherence < 80% -> process recommendation MUST come first.
    if (journal["found"] and journal["rows"]
            and jstats["adherence_pct"] is not None
            and jstats["adherence_pct"] < 80.0):
        recs.insert(0, {
            "text": ("PROCESS FIRST: journal adherence is below the 80% "
                     "gate. Before changing any parameter, run the next sim "
                     "block focused purely on executing the existing plan "
                     "as written (strategy.md section 13)."),
            "evidence": ("Adherence %.1f%% over %d scored journal trades "
                         "(%d chase(s), %d widened stop(s))."
                         % (jstats["adherence_pct"], jstats["adherence_n"],
                            jstats["chases"], jstats["widened"])),
            "confidence": confidence(jstats["adherence_n"]),
        })

    return recs


# ---------------------------------------------------- experiments queue

def build_experiments(records):
    """Concrete queue, smallest-first: each entry names the single variable
    to change, the symbol, and the decision metric."""
    if not records:
        return [
            ("%s -- variable: level_source (PDH/PDL vs ORH/ORL), all other "
             "inputs fixed (Retest only, Fixed 3R, rvol_min 1.2, filters "
             "on); decision metric: higher net R after >= 10 simulated "
             "trades per arm." % sym)
            for sym in WATCHLIST
        ]

    queue = []
    symbols = sorted({r.get("symbol") for r in records if r.get("symbol")})
    by_sym_lvl = aggregate(
        records, lambda r: (r.get("symbol"), r.get("level_source")))

    # 1. Zero-trade symbols: loosen exactly one filter.
    for sym in symbols:
        sym_runs = [r for r in records if r.get("symbol") == sym]
        counts = [_trades(r) for r in sym_runs]
        if counts and all(t == 0 for t in counts):
            rvols = [x for x in (_num(_cfg(r, "rvol_min")) for r in sym_runs)
                     if x is not None]
            cur = min(rvols) if rvols else 1.2
            queue.append(
                "%s -- variable: rvol_min %.1f -> %.1f (nothing else); "
                "decision metric: config produces >= 1 trade/month with "
                "net R >= 0, else drop %s from the active scan."
                % (sym, cur, cur - 0.1, sym))

    # 2. Symbols tested on only one level source: run the other one.
    for sym in symbols:
        have = {lvl for (s, lvl) in by_sym_lvl if s == sym}
        for missing in ("PDH/PDL", "ORH/ORL"):
            if missing not in have:
                queue.append(
                    "%s -- variable: level_source -> %s (config otherwise "
                    "identical to the best existing %s run); decision "
                    "metric: net R difference over >= 10 simulated trades."
                    % (sym, missing, sym))

    # 3. Best net-positive config: vary exit_mode one step.
    scored = [r for r in records
              if _num(_metric(r, "net_r")) is not None
              and (_trades(r) or 0) > 0]
    if scored:
        best = max(scored, key=lambda r: _num(_metric(r, "net_r")))
        cur_exit = _cfg(best, "exit_mode")
        alt = "Trail after 2R" if cur_exit != "Trail after 2R" else "Fixed 3R"
        queue.append(
            "%s -- variable: exit_mode '%s' -> '%s' on the current best "
            "config (%s, %s); decision metric: net R with max DD no worse "
            "than %.2fR." % (best.get("symbol"), cur_exit, alt,
                             best.get("level_source"),
                             _cfg(best, "entry_mode"),
                             _num(_metric(best, "max_dd_r")) or 0.0))

    # 4. Watchlist symbols never backtested at all.
    for sym in WATCHLIST:
        if sym not in symbols:
            queue.append(
                "%s -- variable: first coverage run, level_source PDH/PDL "
                "(baseline config: Retest only, Fixed 3R, rvol_min 1.2); "
                "decision metric: any signal frequency >= 1 trade/month."
                % sym)

    return queue[:8] if queue else build_experiments([])


# ----------------------------------------------------------------- report

def build_session_accuracy_section(session_records, session_warnings,
                                   sessions_path):
    """Section 7: Session decision accuracy -- reads sessions.jsonl (one
    row per trading day, from analysis/record_session.py) and reports the
    confusion matrix, decision accuracy %, QUALIFIED frequency, and total
    pre-window rejects. Prints the ORB same-bar priority warning line on
    every run, regardless of data."""
    lines = []
    add = lines.append
    add("## 7. Session decision accuracy")
    add("")
    add("Source: `%s`. One row per trading day (not per trade), recorded "
        "via `python3 analysis/record_session.py`. Closes the blind spot "
        "where a legitimate no-trade day left zero rows anywhere."
        % os.path.relpath(sessions_path, ROOT))
    add("")

    total = len(session_records)
    if not total:
        add("No sessions recorded yet.")
        add("")
    else:
        matrix = session_confusion_matrix(session_records)
        accuracy = session_decision_accuracy_pct(matrix, total)

        add("### Confusion matrix (engine had a QUALIFIED?) x (trader acted?)")
        add("")
        add("| | Trader acted (trade taken) | Trader did not act |")
        add("|---|---|---|")
        add("| **Engine QUALIFIED (>0)** | CORRECT_TRADE: %d | MISSED_SIGNAL: %d |"
            % (matrix["CORRECT_TRADE"], matrix["MISSED_SIGNAL"]))
        add("| **Engine had no QUALIFIED** | OFF_PLAN_TRADE: %d | CORRECT_NO_TRADE: %d |"
            % (matrix["OFF_PLAN_TRADE"], matrix["CORRECT_NO_TRADE"]))
        add("")
        add("- Sessions recorded: **%d**" % total)
        add("- Decision accuracy: **%s** (target >= %.0f%%, mirroring "
            "strategy.md section 13's adherence gate)"
            % ("n/a" if accuracy is None else "%.1f%%" % accuracy,
               SESSION_ADHERENCE_GATE_PCT))
        if accuracy is not None and accuracy < SESSION_ADHERENCE_GATE_PCT:
            add("  - Below gate. Do not treat this as a parameter problem "
                "first -- it is a process/state-recognition problem "
                "(see MISSED_SIGNAL and OFF_PLAN_TRADE counts above).")
        add("")

        qfreq = session_qualified_frequency(session_records)
        add("### QUALIFIED frequency")
        add("")
        if qfreq["sessions_with_known_engine"]:
            add("- QUALIFIED events per session: **%.2f** (%d total across "
                "%d session(s) with a known engine replay)"
                % (qfreq["per_session"], qfreq["total_qualified"],
                   qfreq["sessions_with_known_engine"]))
            add("- Approx. QUALIFIED events per week (x%d trading days): "
                "**%.2f**" % (TRADING_DAYS_PER_WEEK, qfreq["per_week"]))
            add("  - If this stays near zero as more sessions accumulate, "
                "the edge is unvalidatable at this frequency no matter how "
                "disciplined execution is -- that is a separate problem "
                "from adherence.")
        else:
            add("- No sessions with a known engine replay yet (all "
                "recorded via --no-events).")
        if qfreq["excluded"]:
            add("  - %d session(s) excluded from this rate (recorded with "
                "--no-events; engine QUALIFIED count unknown, not assumed "
                "zero)." % qfreq["excluded"])
        add("")

        pw_total, pw_known, pw_excluded = session_total_pre_window_rejects(
            session_records)
        add("### Pre-window rejects")
        add("")
        add("- Total pre-window REJECTs (signal before the 08:45 CT entry "
            "window opened): **%d** across %d session(s) with a known count."
            % (pw_total, pw_known))
        if pw_excluded:
            add("  - %d session(s) excluded (engine count unknown)."
                % pw_excluded)
        add("  - Open question (see daytrading_memory.md / NEXT_SESSION_PLAN): "
            "should a pre-window REJECT burn the track for the rest of the "
            "day, or should it re-arm? Not enough sessions yet to decide.")
        add("")

    if session_warnings:
        add("Appendix note: %d issue(s) reading sessions.jsonl (see the "
            "appendix below)." % len(session_warnings))
        add("")

    add("**%s**" % ORB_TIEBREAK_WARNING)
    add("")
    return lines


def build_report(records, warnings, journal, results_path, journal_path,
                 session_records=None, session_warnings=None,
                 sessions_path=DEFAULT_SESSIONS):
    jrows = journal["rows"]
    jstats = journal_stats(jrows) if jrows else None
    recs = build_recommendations(records, journal, jstats or
                                 {"adherence_pct": None})
    experiments = build_experiments(records)

    symbols = sorted({r.get("symbol") for r in records if r.get("symbol")})
    configs = {config_signature(r) for r in records}
    trade_counts = [_trades(r) for r in records if _trades(r) is not None]
    total_trades = sum(trade_counts) if trade_counts else 0

    lines = []
    add = lines.append
    add("# Learning-Loop Report")
    add("")
    add("Generated: %s CT | results: `%s` | journal: `%s`"
        % (datetime.now().strftime("%Y-%m-%d %H:%M"),
           os.path.relpath(results_path, ROOT),
           os.path.relpath(journal_path, ROOT)
           if os.path.exists(journal_path) else journal_path))
    add("")
    add("ADVISORY ONLY -- decision support for simulation planning. "
        "No order flow, no brokerage interaction.")
    add("")

    # 1. Data inventory --------------------------------------------------
    add("## 1. Data inventory")
    add("")
    add("- Backtest runs recorded: **%d**" % len(records))
    add("- Symbols covered: %s"
        % (", ".join(symbols) if symbols else "none"))
    add("- Distinct configs covered: %d" % len(configs))
    add("- Total simulated trades: %d" % total_trades)
    if journal["found"]:
        live = sum(1 for r in jrows if not r["is_sim"])
        sim = sum(1 for r in jrows if r["is_sim"])
        add("- Journal rows found: %d (live: %d, sim: %d)"
            % (len(jrows), live, sim))
        if journal["examples_excluded"]:
            add("  - %d worked-example row(s) detected and excluded "
                "(Screenshot = `%s`)."
                % (journal["examples_excluded"], EXAMPLE_SCREENSHOT))
        if not jrows:
            add("  - The journal has no real trade rows yet -- only the "
                "worked example.")
    else:
        add("- Journal: %s" % (journal["note"] or "not read"))
    add("")

    # 2. Backtest league table -------------------------------------------
    add("## 2. Backtest league table")
    add("")
    if records:
        add("Sorted by net R (desc). Runs with < 10 trades are flagged "
            "as tiny samples.")
        add("")
        add("| run_id | symbol | level | entry | exit | trades | net R | "
            "win % | PF | max DD (R) | flags |")
        add("|---|---|---|---|---|---|---|---|---|---|---|")

        def sort_key(r):
            nr = _num(_metric(r, "net_r"))
            return -(nr if nr is not None else float("-inf")), str(r.get("run_id"))

        for r in sorted(records, key=sort_key):
            t = _trades(r)
            flags = ("tiny sample" if (t is None or t < 10) else "")
            add("| %s | %s | %s | %s | %s | %s | %s | %s | %s | %s | %s |" % (
                r.get("run_id"), r.get("symbol"), r.get("level_source"),
                _cfg(r, "entry_mode"), _cfg(r, "exit_mode"),
                "n/a" if t is None else t,
                _fmt(_num(_metric(r, "net_r"))),
                _fmt(_num(_metric(r, "win_rate_pct")), 1),
                _fmt(_num(_metric(r, "profit_factor_r"))),
                _fmt(_num(_metric(r, "max_dd_r"))),
                flags))
    else:
        add("No backtest runs recorded yet. Record runs with "
            "`python3 analysis/record_backtest.py --json '{...}'`.")
    add("")

    # 3. Aggregations -----------------------------------------------------
    add("## 3. Aggregations (dimensions with >= 2 runs)")
    add("")
    dims = (
        ("By symbol", lambda r: r.get("symbol")),
        ("By level source", lambda r: r.get("level_source")),
        ("By entry mode", lambda r: _cfg(r, "entry_mode")),
        ("By exit mode", lambda r: _cfg(r, "exit_mode")),
    )
    any_agg = False
    for title, keyfunc in dims:
        groups = {k: g for k, g in aggregate(records, keyfunc).items()
                  if g["runs"] >= 2 and k is not None}
        if not groups:
            continue
        any_agg = True
        add("### %s" % title)
        add("")
        add("| group | runs | total trades | combined net R | "
            "weighted win % |")
        add("|---|---|---|---|---|")
        for k in sorted(groups, key=str):
            g = groups[k]
            add("| %s | %d | %s | %s | %s |" % (
                k, g["runs"],
                "n/a" if g["trades"] is None else g["trades"],
                _fmt(g["net_r"]),
                _fmt(weighted_win_rate(g), 1)))
        add("")
    if not any_agg:
        add("No dimension has >= 2 runs yet; aggregations will appear as "
            "the results file grows.")
        add("")

    # 4. Journal-based execution metrics ----------------------------------
    add("## 4. Journal-based execution metrics")
    add("")
    if not journal["found"]:
        add("Journal unavailable: %s. Section skipped; the rest of the "
            "report is unaffected." % (journal["note"] or "not read"))
    elif not jrows:
        add("No real journal rows yet (only the worked-example row, which "
            "was detected and excluded). Execution metrics will populate "
            "once real live/sim trades are journaled.")
    else:
        add("- Trades journaled: %d" % jstats["n"])
        add("- Adherence rate: %s (over %d scored trades)"
            % ("n/a" if jstats["adherence_pct"] is None
               else "%.1f%%" % jstats["adherence_pct"],
               jstats["adherence_n"]))
        add("- Chase count: %d" % jstats["chases"])
        add("- Widened-stop count: %d" % jstats["widened"])
        if jstats["r_by_setup"]:
            add("- Realized R by setup type:")
            for setup, (avg, tot, n) in jstats["r_by_setup"].items():
                add("  - Setup %s: avg %.2fR, total %.2fR over %d trade(s)"
                    % (setup, avg, tot, n))
        add("- Average MFE/MAE utilization: %s"
            % ("n/a (needs Realized R + Highest chart R)"
               if jstats["mfe_util_pct"] is None
               else "%.0f%% of best available chart R captured"
               % jstats["mfe_util_pct"]))
        if jstats["avg_mfe"] is not None or jstats["avg_mae"] is not None:
            add("  - Avg MFE: %s | Avg MAE: %s (option $ terms)"
                % (_fmt(jstats["avg_mfe"]), _fmt(jstats["avg_mae"])))
        if jstats["followed_r"] and jstats["broken_r"]:
            add("- Correlation callout (association, not causation): in %d "
                "trades where plan was not followed, avg R was %.2f vs "
                "%.2f in the %d plan-followed trades."
                % (len(jstats["broken_r"]),
                   sum(jstats["broken_r"]) / len(jstats["broken_r"]),
                   sum(jstats["followed_r"]) / len(jstats["followed_r"]),
                   len(jstats["followed_r"])))
    add("")

    # 5. Recommendations ---------------------------------------------------
    add("## 5. Recommendations")
    add("")
    add("Confidence is LOW unless the supporting sample is >= 20 trades, "
        "and never above MEDIUM from backtests alone.")
    add("")
    if recs:
        for i, rec in enumerate(recs, 1):
            add("%d. **[%s]** %s" % (i, rec["confidence"], rec["text"]))
            add("   - Evidence: %s" % rec["evidence"])
    else:
        add("No rule fired on the current evidence. Keep collecting runs "
            "and journal rows; the experiments queue below is the path to "
            "more evidence.")
    add("")

    # 6. Next experiments queue --------------------------------------------
    add("## 6. Next experiments queue (smallest first)")
    add("")
    add("Each experiment changes exactly one variable and names its "
        "decision metric. Run in Bar Replay / strategy tester only.")
    add("")
    for i, exp in enumerate(experiments, 1):
        add("%d. %s" % (i, exp))
    add("")

    # 7. Session decision accuracy ------------------------------------------
    lines.extend(build_session_accuracy_section(
        session_records or [], session_warnings or [], sessions_path))

    # Appendix -------------------------------------------------------------
    all_warnings = list(warnings) + list(session_warnings or [])
    if all_warnings:
        add("## Appendix: skipped/malformed input")
        add("")
        for w in all_warnings:
            add("- %s" % w)
        add("")

    add("---")
    add(FOOTER)
    add("")
    return "\n".join(lines)


def generate(results_path, journal_path, out_path, sessions_path=DEFAULT_SESSIONS):
    records, warnings = load_results(results_path)
    journal = load_journal(journal_path)
    session_records, session_warnings = load_sessions(sessions_path)
    report = build_report(records, warnings, journal, results_path,
                          journal_path, session_records=session_records,
                          session_warnings=session_warnings,
                          sessions_path=sessions_path)
    out_dir = os.path.dirname(os.path.abspath(out_path))
    if out_dir and not os.path.isdir(out_dir):
        os.makedirs(out_dir)
    with open(out_path, "w") as fh:
        fh.write(report)
    return report


def main(argv=None):
    parser = argparse.ArgumentParser(
        description="Learning-loop engine: backtests + journal -> advisory "
                    "report (no brokerage interaction).")
    parser.add_argument("--results", default=DEFAULT_RESULTS)
    parser.add_argument("--journal", default=DEFAULT_JOURNAL)
    parser.add_argument("--sessions", default=DEFAULT_SESSIONS)
    parser.add_argument("--out", default=DEFAULT_OUT)
    args = parser.parse_args(argv)
    generate(args.results, args.journal, args.out, args.sessions)
    print("report written -> %s" % args.out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
